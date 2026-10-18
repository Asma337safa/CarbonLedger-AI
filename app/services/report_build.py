"""Assemblage des rapports (PDF de gestion, PDF de décision, export Excel).

Logique partagée entre l'interface web FastAPI et l'application Streamlit.
Aucun chiffre n'est calculé ici : montants et émissions proviennent de
``app.calc`` (mêmes fonctions que le tableau de bord), et la mise en page de
``app.services.report_pdf``.
"""

from __future__ import annotations

import io
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session as OrmSession

from ..calc.dashboard import category_breakdown, emissions_breakdown, summarize_expenses
from ..config import settings
from ..enums import CARBON_CATEGORIES, EXPENSE_CATEGORIES, PRIORITIES, PROJECT_STATUSES
from ..errors import ValidationError
from ..models import CarbonRecord, DecisionProject, Expense, Report
from . import audit
from .project_service import analyze_project, project_inputs
from .report_pdf import generate_decision_pdf, generate_period_pdf


def collect_period_data(db: OrmSession, organization_id: int, resolved):
    """Dépenses et enregistrements carbone de la période résolue."""
    expenses = (
        db.execute(
            select(Expense).where(
                Expense.organization_id == organization_id,
                Expense.expense_date >= resolved.start,
                Expense.expense_date <= resolved.end,
            )
        )
        .scalars()
        .all()
    )
    records = (
        db.execute(
            select(CarbonRecord).where(
                CarbonRecord.organization_id == organization_id,
                CarbonRecord.period_start <= resolved.end,
                CarbonRecord.period_end >= resolved.start,
            )
        )
        .scalars()
        .all()
    )
    return expenses, records


def factor_sources(db: OrmSession, organization_id: int, records) -> list[str]:
    """Sources des facteurs d'émission utilisés (traçabilité du rapport)."""
    sources: list[str] = []
    seen: set[str] = set()
    for record in records:
        factor = record.emission_factor
        if factor is None:
            continue
        text = f"{factor.code} — {factor.label} : {factor.kgco2e_per_unit} kgCO₂e/{factor.unit} — source : {factor.source_name or 'non documentée'}"
        if text not in seen:
            seen.add(text)
            sources.append(text)
    return sources


def period_notes(resolved, expense_summary) -> list[str]:
    """Notes de méthode du rapport de gestion (transparence des limites)."""
    return [
        f"Période analysée : {resolved.label}.",
        "Les montants proviennent des dépenses saisies ou importées, hors éléments non renseignés.",
        (
            f"{expense_summary.without_factor_count} dépense(s) sans facteur d'émission : "
            "leur carbone n'est pas estimé (aucun facteur attribué automatiquement)."
        ),
    ]


def projects_rows(db: OrmSession, organization_id: int) -> list[dict[str, Any]]:
    """Synthèse des projets de décision (section du rapport de gestion)."""
    projects = (
        db.execute(
            select(DecisionProject).where(DecisionProject.organization_id == organization_id)
        )
        .scalars()
        .all()
    )
    return [
        {
            "title": project.title,
            "status": PROJECT_STATUSES.get(project.status, project.status),
            "summary": (project.description or "")[:200],
        }
        for project in projects
    ]


def generate_period_report(
    db: OrmSession,
    *,
    organization,
    user,
    resolved,
    request=None,
) -> Report:
    """Crée le rapport PDF de gestion d'une période (fichier + entrée Report)."""
    expenses, records = collect_period_data(db, organization.id, resolved)
    expense_summary = summarize_expenses(expenses)
    emission_summary, emission_rows = emissions_breakdown(records)
    category_rows = category_breakdown(expenses)
    sources = factor_sources(db, organization.id, records)
    notes = period_notes(resolved, expense_summary)

    target_dir = settings.reports_dir / str(organization.id)
    target_dir.mkdir(parents=True, exist_ok=True)
    # Pour « Tout l'historique », on historise les bornes réelles des données
    # (sinon le journal afficherait 01/01/2000 → 31/12/9999).
    period_start, period_end = resolved.start, resolved.end
    if resolved.key == "tout":
        dates = [row.expense_date for row in expenses if row.expense_date]
        if dates:
            period_start, period_end = min(dates), max(dates)
    report = Report(
        organization_id=organization.id,
        kind="pdf",
        scope="periode",
        title=f"Rapport de gestion — {resolved.label}",
        period_start=period_start,
        period_end=period_end,
        created_by_id=user.id,
    )
    db.add(report)
    db.flush()
    path = target_dir / f"rapport_gestion_{report.id}.pdf"
    report.payload = generate_period_pdf(
        path,
        organization_name=organization.name,
        period_label=resolved.label,
        expense_summary=expense_summary,
        category_rows=category_rows,
        emission_summary=emission_summary,
        emission_rows=emission_rows,
        projects_rows=projects_rows(db, organization.id),
        factor_sources=sources,
        notes=notes,
        author_name=user.full_name or user.email,
    )
    report.file_path = str(path)
    audit.log(
        db,
        action="report_generated",
        organization_id=organization.id,
        user_id=user.id,
        entity_type="report",
        entity_id=report.id,
        details={"scope": "periode", "periode": resolved.label},
        request=request,
    )
    db.commit()
    return report


def generate_project_report(
    db: OrmSession,
    *,
    organization,
    user,
    project,
    request=None,
) -> Report:
    """Crée le rapport PDF d'un projet de décision (fichier + entrée Report)."""
    inputs = project_inputs(project)
    if len(inputs) < 2:
        raise ValidationError("Ce projet comporte moins de deux options : rapport impossible.")
    analysis = analyze_project(project)

    target_dir = settings.reports_dir / str(organization.id)
    target_dir.mkdir(parents=True, exist_ok=True)
    report = Report(
        organization_id=organization.id,
        project_id=project.id,
        kind="pdf",
        scope="projet",
        title=f"Rapport de décision — {project.title}",
        created_by_id=user.id,
    )
    db.add(report)
    db.flush()
    path = target_dir / f"rapport_decision_{project.id}_{report.id}.pdf"
    report.payload = generate_decision_pdf(
        path,
        organization_name=organization.name,
        project_title=project.title,
        project_category=EXPENSE_CATEGORIES.get(project.category, project.category),
        horizon_years=project.horizon_years,
        priority_label=PRIORITIES.get(project.priority, project.priority),
        carbon_price=project.internal_carbon_price,
        analysis=analysis,
        author_name=user.full_name or user.email,
    )
    report.file_path = str(path)
    audit.log(
        db,
        action="report_generated",
        organization_id=organization.id,
        user_id=user.id,
        entity_type="report",
        entity_id=report.id,
        details={"scope": "projet", "projet": project.title},
        request=request,
    )
    db.commit()
    return report


def export_period_excel(db: OrmSession, *, organization_id: int, resolved) -> tuple[str, bytes]:
    """Export Excel de la période : dépenses, registre carbone, décisions.

    Renvoie ``(nom de fichier, contenu .xlsx)``.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    expenses, records = collect_period_data(db, organization_id, resolved)
    projects = (
        db.execute(
            select(DecisionProject).where(DecisionProject.organization_id == organization_id)
        )
        .scalars()
        .all()
    )

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="176B4D")

    sheet = workbook.active
    sheet.title = "Dépenses"
    sheet.append([
        "Date", "Catégorie", "Fournisseur", "Description", "Montant HT", "TVA", "Montant TTC",
        "Quantité", "Unité", "Facteur", "Statut", "Démonstration",
    ])
    for expense in expenses:
        sheet.append([
            expense.expense_date.isoformat(),
            EXPENSE_CATEGORIES.get(expense.category, expense.category),
            expense.supplier,
            expense.description,
            float(expense.amount_ht),
            float(expense.vat_amount),
            float(expense.amount_ttc),
            float(expense.quantity) if expense.quantity is not None else None,
            expense.unit or "",
            expense.emission_factor.code if expense.emission_factor else "",
            expense.verification_status,
            "oui" if expense.is_demo else "non",
        ])

    sheet = workbook.create_sheet("Registre carbone")
    sheet.append([
        "Début", "Fin", "Catégorie", "Quantité d'activité", "Unité", "Facteur", "kg CO2e", "t CO2e",
        "Méthode", "Statut", "Confiance", "Source du facteur",
    ])
    for record in records:
        sheet.append([
            record.period_start.isoformat(),
            record.period_end.isoformat(),
            CARBON_CATEGORIES.get(record.category, record.category),
            float(record.activity_quantity),
            record.activity_unit,
            record.emission_factor.code if record.emission_factor else "",
            float(record.kgco2e),
            float(record.kgco2e / 1000),
            record.method,
            record.status,
            record.confidence,
            record.emission_factor.source_name if record.emission_factor else "",
        ])

    sheet = workbook.create_sheet("Décisions")
    sheet.append([
        "Projet", "Statut", "Option", "Fournisseur", "Coût d'acquisition", "Coût total (horizon)",
        "Émissions cumulées (t)", "Confiance", "Options complètes",
    ])
    for project in projects:
        inputs = project_inputs(project)
        if len(inputs) < 2:
            sheet.append([project.title, project.status, "—", "", None, None, None, "", "moins de deux options"])
            continue
        try:
            from ..calc.decision import compare_options

            comparison = compare_options(inputs, project.horizon_years, carbon_price=project.internal_carbon_price)
            for result in comparison.results:
                sheet.append([
                    project.title,
                    PROJECT_STATUSES.get(project.status, project.status),
                    result.name,
                    result.option.supplier,
                    float(result.purchase_cost),
                    float(result.total_cost),
                    float(result.total_emissions_t) if result.total_emissions_t is not None else None,
                    result.option.confidence,
                    "",
                ])
        except Exception:
            sheet.append([project.title, project.status, "—", "", None, None, None, "", "calcul impossible"])

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for column_cells in sheet.columns:
            length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(42, max(12, length + 2))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return f"carbonledger_export_{resolved.range_slug}.xlsx", buffer.getvalue()
